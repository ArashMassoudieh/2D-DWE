#!/usr/bin/env python3
"""
Point shapefile of the WS3 culvert field measurements, straight from the form.

    python3 tools/culvert_form_to_shp.py

Coordinates come from the form's LOCATION block, which the form itself marks
"pre-filled, please do not edit" -- so the ID-to-location mapping is the
office's, set before the site visit, and is the authoritative link between a
row of measurements and a place on the ground.
"""
import os, re
import openpyxl
from osgeo import ogr, osr
ogr.UseExceptions()

XL  = "/home/arash/Dropbox/Chloeta/NHA/WS3_Culvert_Field_Measurement_Form (003).xlsx"
OUT = "/home/arash/Dropbox/Chloeta/NHA/WS3_GIS/NM15-141Local/point/WS3_culverts_measured.shp"

wb = openpyxl.load_workbook(XL, data_only=True); ws = wb["Field Form"]
hr = next(r for r in range(1, 12)
          if any(str(c.value).strip() == "ID" for c in ws[r] if c.value))
hdr = [re.sub(r"\s+", " ", str(c.value).strip()) if c.value else "" for c in ws[hr]]

# shapefile field names are capped at 10 characters
FIELDS = [("ID", "id", ogr.OFTInteger), ("Cluster", "cluster", ogr.OFTString),
          ("Project No.", "project", ogr.OFTString), ("Community", "community", ogr.OFTString),
          ("Road / Highway", "road", ogr.OFTString), ("Watercourse", "watercours", ogr.OFTString),
          ("Structure present?", "present", ogr.OFTString), ("Culvert shape", "shape", ogr.OFTString),
          ("No. of barrels", "barrels", ogr.OFTString), ("Diameter or Span (in)", "span_in", ogr.OFTString),
          ("Rise / Height (in)", "rise_in", ogr.OFTString), ("Barrel length (ft)", "length_ft", ogr.OFTString),
          ("Material", "material", ogr.OFTString), ("Inlet end treatment", "inlet_trt", ogr.OFTString),
          ("Outlet end treatment", "outlet_trt", ogr.OFTString), ("Skew to road (deg)", "skew", ogr.OFTString),
          ("Road surface to INLET invert (ft)", "fill_in", ogr.OFTString),
          ("Road surface to OUTLET invert (ft)", "fill_out", ogr.OFTString),
          ("Channel width u/s (ft)", "chan_w", ogr.OFTString),
          ("Channel depth u/s (ft)", "chan_d", ogr.OFTString),
          ("Blockage / siltation (%)", "blockage", ogr.OFTString),
          ("Overall condition", "condition", ogr.OFTString),
          ("Condition notes", "cond_note", ogr.OFTString),
          ("Site notes", "site_note", ogr.OFTString),
          ("Photos taken?", "photos", ogr.OFTString),
          ("Photo file refs", "photo_ref", ogr.OFTString),
          ("Measured by", "meas_by", ogr.OFTString), ("Date measured", "meas_date", ogr.OFTString)]
idx = {}
for want, short, typ in FIELDS:
    key = re.sub(r"\s+", " ", want)
    idx[short] = next((i for i, h in enumerate(hdr) if re.sub(r"\s+", " ", h) == key), None)
ilat = next(i for i, h in enumerate(hdr) if h == "Latitude")
ilon = next(i for i, h in enumerate(hdr) if h == "Longitude")

drv = ogr.GetDriverByName("ESRI Shapefile")
if os.path.exists(OUT): drv.DeleteDataSource(OUT)
ds = drv.CreateDataSource(OUT)
srs = osr.SpatialReference(); srs.ImportFromEPSG(4326)
ly = ds.CreateLayer("WS3_culverts_measured", srs, ogr.wkbPoint)
for want, short, typ in FIELDS:
    fd = ogr.FieldDefn(short, typ)
    if typ == ogr.OFTString: fd.SetWidth(120)
    ly.CreateField(fd)
ly.CreateField(ogr.FieldDefn("lat", ogr.OFTReal)); ly.CreateField(ogr.FieldDefn("lon", ogr.OFTReal))

n = 0
for row in ws.iter_rows(min_row=hr + 1, max_row=ws.max_row):
    rid = row[0].value
    if rid is None or str(rid).strip() in ("", "EX"): continue
    try: lat = float(row[ilat].value); lon = float(row[ilon].value)
    except (TypeError, ValueError): continue
    f = ogr.Feature(ly.GetLayerDefn())
    for want, short, typ in FIELDS:
        i = idx[short]
        if i is None or i >= len(row): continue
        v = row[i].value
        if v in (None, ""): continue
        if typ == ogr.OFTInteger:
            try: f.SetField(short, int(v))
            except (TypeError, ValueError): pass
        else:
            f.SetField(short, re.sub(r"\s+", " ", str(v)).strip()[:120])
    f.SetField("lat", lat); f.SetField("lon", lon)
    g = ogr.Geometry(ogr.wkbPoint); g.AddPoint(lon, lat); f.SetGeometry(g)
    ly.CreateFeature(f); n += 1
ds = None
print("wrote %s  (%d culverts, EPSG:4326)" % (OUT, n))
