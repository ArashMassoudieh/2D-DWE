#!/usr/bin/env python3
"""
Build the C10 crossing-sizing workbook. Every derived cell is a live Excel
formula; only the Inputs sheet and the surveyed cross-section hold numbers.
"""
import csv, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SP = "/tmp/claude-1000/-home-arash-Projects-2D-DWE/615dd2c4-6a45-4d7c-8d78-98fee05c9920/scratchpad"
OUT = "/home/arash/Dropbox/Chloeta/NHA/WS3_GIS/NM15-141Local/Report_for_C10/C10_crossing_sizing.xlsx"

HDR = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E79")
SUB = Font(bold=True); INP = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E2EFDA"); NOTE = Font(italic=True, size=9, color="555555")
BOX = Border(*[Side(style="thin", color="BBBBBB")] * 4)

wb = openpyxl.Workbook()

# ------------------------------------------------------------------ Inputs --
ws = wb.active; ws.title = "Inputs"
ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 10; ws.column_dimensions["D"].width = 62
ws["A1"] = "C10 / NM15-141 — Navajo Route 9 crossing: sizing inputs"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Yellow cells are inputs. Everything on the other sheets is a formula."; ws["A2"].font = NOTE
r = 4
for lbl, key, val, unit, note in [
    ("HYDROLOGY", None, None, None, None),
    ("Design discharge Q100 (6-hr, converged HEC-HMS)", "Q", 11922, "cfs", "converged Tc-Q iteration; pass-1 value was 10065"),
    ("GEOMETRY (1 m LiDAR)", None, None, None, None),
    ("Inlet invert", "INV_IN", 6112.80, "ft", "upstream face"),
    ("Outlet invert", "INV_OUT", 6110.60, "ft", "downstream face"),
    ("Road low point", "ROAD", 6120.60, "ft", "controls allowable headwater"),
    ("Barrel length", "L", 130.0, "ft", "through the embankment"),
    ("Freeboard below road", "FB", 1.00, "ft", "set 0 to size to the road itself"),
    ("PHYSICAL", None, None, None, None),
    ("Gravity g", "G", 32.174, "ft/s2", ""),
    ("Weir coefficient (broad-crested, check only)", "CW", 3.00, "-", "used only as an independent check"),
    ("Discharge coefficient (orifice check)", "CD", 0.60, "-", "used only for the existing-barrel check"),
    ("TAILWATER", None, None, None, None),
    ("Downstream bed slope S0_ds (tracked thalweg)", "S_DS", 0.0121, "ft/ft", "27.4 ft over 2264 ft"),
    ("Channel Manning n (tailwater)", "N_CH", 0.045, "-", "NLCD-derived range 0.035-0.050"),
]:
    if key is None:
        ws.cell(r, 1, lbl).font = SUB; r += 1; continue
    ws.cell(r, 1, lbl); c = ws.cell(r, 2, val); c.fill = INP; c.border = BOX
    ws.cell(r, 3, unit); ws.cell(r, 4, note).font = NOTE
    wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName(key, attr_text="Inputs!$B$%d" % r))
    r += 1

r += 1
ws.cell(r, 1, "ENTRANCE COEFFICIENTS (FHWA HDS-5)").font = SUB; r += 1
hrow = r
for j, h in enumerate(["Entrance type", "shape", "span/dia (ft)", "rise (ft)", "c", "Y", "ke", "n barrel"], 1):
    c = ws.cell(r, j, h); c.font = HDR; c.fill = HF
r += 1
ENT = [("CMP, square edge with headwall", "circ", 6, 6, 0.0379, 0.69, 0.50, 0.024),
       ("CMP, mitered to slope",          "circ", 6, 6, 0.0463, 0.75, 0.70, 0.024),
       ("RCP, groove end with headwall",  "circ", 6, 6, 0.0292, 0.74, 0.20, 0.012),
       ("RCB, 90 deg headwall chamfered", "box",  8, 6, 0.0347, 0.81, 0.50, 0.012),
       ("RCB, 90 deg headwall chamfered", "box", 12,10, 0.0347, 0.81, 0.50, 0.012)]
ENT_R0 = r
for e in ENT:
    for j, v in enumerate(e, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if j >= 3: c.fill = INP
    r += 1
ENT_R1 = r - 1
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("ENT", attr_text="Inputs!$A$%d:$H$%d" % (ENT_R0, ENT_R1)))

r += 1
ws.cell(r, 1, "EXISTING STRUCTURE (field measured 17 Aug 2026)").font = SUB; r += 1
for j, h in enumerate(["Description", "barrels", "diameter (in)"], 1):
    c = ws.cell(r, j, h); c.font = HDR; c.fill = HF
r += 1
EX_R0 = r
for d in [("culvert 8 as tabulated", 2, 40), ("culvert 8 as photographed", 2, 85)]:
    for j, v in enumerate(d, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if j >= 2: c.fill = INP
    r += 1
EX_R1 = r - 1

# ------------------------------------------------------- derived quantities --
r += 1
ws.cell(r, 1, "DERIVED (formulas)").font = SUB; r += 1
for lbl, key, f, unit in [
    ("Barrel slope S0 = (inlet inv - outlet inv)/L", "S0", "=(INV_IN-INV_OUT)/L", "ft/ft"),
    ("Available fill (road - inlet invert)", "FILL", "=ROAD-INV_IN", "ft"),
    ("Allowable headwater HWmax", "HWMAX", "=ROAD-INV_IN-FB", "ft"),
]:
    ws.cell(r, 1, lbl); c = ws.cell(r, 2, f); c.fill = CALC; c.border = BOX
    ws.cell(r, 3, unit)
    wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName(key, attr_text="Inputs!$B$%d" % r))
    r += 1

# ----------------------------------------------------------- outlet section --
sec = wb.create_sheet("Section_Outlet")
sec["A1"] = "Surveyed outlet cross-section (1 m LiDAR, y = 3,960,500, every 6.56 ft)"
sec["A1"].font = Font(bold=True, size=12)
sec["A2"] = "The only other numbers in this workbook. Everything downstream of here is formulas."
sec["A2"].font = NOTE
for j, h in enumerate(["offset (ft)", "bed elevation (ft)"], 1):
    c = sec.cell(4, j, h); c.font = HDR; c.fill = HF
rows = list(csv.DictReader(open(SP + "/sec_outlet.csv")))
for i, d in enumerate(rows):
    sec.cell(5 + i, 1, float(d["offset_ft"])).fill = INP
    sec.cell(5 + i, 2, float(d["elev_ft"])).fill = INP
n = len(rows); LAST = 4 + n
sec.column_dimensions["A"].width = 13; sec.column_dimensions["B"].width = 19
sec["D4"] = "station spacing dx (ft)"; sec["D5"] = "=Section_Outlet!A6-Section_Outlet!A5"
sec["D6"] = "invert (ft)"; sec["D7"] = "=MIN(Section_Outlet!B5:B%d)" % LAST
for cc in ("D5", "D7"): sec[cc].fill = CALC; sec[cc].border = BOX
sec.column_dimensions["D"].width = 24
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("SEC_Z", attr_text="Section_Outlet!$B$5:$B$%d" % LAST))
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("DX", attr_text="Section_Outlet!$D$5"))
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("BED", attr_text="Section_Outlet!$D$7"))

# ---------------------------------------------------------------- tailwater --
tw = wb.create_sheet("Tailwater")
tw["A1"] = "Tailwater at the outlet — stage table, critical depth and normal depth"
tw["A1"].font = Font(bold=True, size=12)
tw["A2"] = ("A(d) and T(d) are computed straight off the surveyed section for each trial depth, "
            "so no channel shape is assumed."); tw["A2"].font = NOTE
hdrs = ["depth d (ft)", "WSE (ft)", "area A (sf)", "top width T (ft)",
        "wetted P (ft)", "R = A/P (ft)", "Froude fn  Q^2*T/(g*A^3)",
        "Manning Q (cfs)"]
for j, h in enumerate(hdrs, 1):
    c = tw.cell(4, j, h); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    tw.column_dimensions[get_column_letter(j)].width = 17
NSTEP = 80
for i in range(NSTEP):
    R = 5 + i
    tw.cell(R, 1, "=%s" % ("0.25" if i == 0 else "A%d+0.25" % (R - 1)))
    tw.cell(R, 2, "=BED+A%d" % R)
    tw.cell(R, 3, "=SUMPRODUCT((B%d-SEC_Z)*(B%d>SEC_Z))*DX" % (R, R))
    tw.cell(R, 4, "=SUMPRODUCT(--(B%d>SEC_Z))*DX" % R)
    tw.cell(R, 5, "=D%d+2*A%d" % (R, R))
    tw.cell(R, 6, "=IF(E%d>0,C%d/E%d,0)" % (R, R, R))
    tw.cell(R, 7, "=IF(C%d>0,Q^2*D%d/(G*C%d^3),\"\")" % (R, R, R))
    tw.cell(R, 8, "=IF(C%d>0,(1.49/N_CH)*C%d*F%d^(2/3)*SQRT(S_DS),0)" % (R, R, R))
    for j in range(1, 9): tw.cell(R, j).fill = CALC
LASTT = 4 + NSTEP
b = LASTT + 2
tw.cell(b, 1, "CRITICAL DEPTH (Froude function = 1)").font = SUB
tw.cell(b + 1, 1, "row below the crossing")
tw.cell(b + 1, 2, "=MATCH(1,G5:G%d,-1)" % LASTT)
tw.cell(b + 2, 1, "critical depth dc (ft)")
tw.cell(b + 2, 2, "=INDEX(A5:A%d,B%d)+0.25*(INDEX(G5:G%d,B%d)-1)/(INDEX(G5:G%d,B%d)-INDEX(G5:G%d,B%d+1))"
        % (LASTT, b + 1, LASTT, b + 1, LASTT, b + 1, LASTT, b + 1))
tw.cell(b + 3, 1, "critical WSE (ft)"); tw.cell(b + 3, 2, "=BED+B%d" % (b + 2))
tw.cell(b + 4, 1, "TW above inlet invert (ft)"); tw.cell(b + 4, 2, "=B%d-INV_IN" % (b + 3))
tw.cell(b + 5, 1, "TW above outlet invert = ho (ft)"); tw.cell(b + 5, 2, "=B%d-INV_OUT" % (b + 3))
tw.cell(b + 7, 1, "NORMAL DEPTH (Manning Q = Q100), for comparison").font = SUB
tw.cell(b + 8, 1, "row"); tw.cell(b + 8, 2, "=MATCH(Q,H5:H%d,1)" % LASTT)
tw.cell(b + 9, 1, "normal depth (ft)")
# interpolate between the bracketing 0.25 ft rows, as the critical-depth block above does.
# MATCH(...,1) alone returns the grid value BELOW the answer, which truncated d_n by up to
# 0.25 ft (it reported 9.50 ft where the true normal depth is 9.69 ft) and so understated
# the governing tailwater and every outlet-control headwater derived from it.
tw.cell(b + 9, 2, "=INDEX(A5:A{L},B{r})+0.25*(Q-INDEX(H5:H{L},B{r}))"
                  "/(INDEX(H5:H{L},B{r}+1)-INDEX(H5:H{L},B{r}))".format(L=LASTT, r=b + 8))
tw.cell(b + 10, 1, "normal WSE (ft)"); tw.cell(b + 10, 2, "=BED+B%d" % (b + 9))
tw.cell(b + 12, 1, "GOVERNING TAILWATER = the higher of critical and normal").font = SUB
tw.cell(b + 13, 1, "ho used in outlet control (ft above outlet invert)")
tw.cell(b + 13, 2, "=MAX(B%d,B%d-INV_OUT)" % (b + 5, b + 10))
for rr in (b + 1, b + 2, b + 3, b + 4, b + 5, b + 8, b + 9, b + 10, b + 13):
    tw.cell(rr, 2).fill = CALC; tw.cell(rr, 2).border = BOX
tw.column_dimensions["A"].width = 46
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("HO", attr_text="Tailwater!$B$%d" % (b + 13)))
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("TW_IN", attr_text="Tailwater!$B$%d" % (b + 4)))

# ------------------------------------------------------------ inlet control --
ic = wb.create_sheet("InletControl")
ic["A1"] = "Inlet control"; ic["A1"].font = Font(bold=True, size=12)
ic["A2"] = ("Unsubmerged (critical flow at the entrance) is the valid regime for a wide opening at this head; "
            "the HDS-5 submerged form is shown with its validity flag."); ic["A2"].font = NOTE
ic["A4"] = "A — UNSUBMERGED: critical flow, per foot of opening"; ic["A4"].font = SUB
for j, h in enumerate(["HW (ft)", "dc = HW/1.5 (ft)", "q = SQRT(g*dc^3) (cfs/ft)",
                       "weir check CW*HW^1.5", "opening for Q100 (ft)", "inlet V (ft/s)"], 1):
    c = ic.cell(5, j, h); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    ic.column_dimensions[get_column_letter(j)].width = 19
for i, hw in enumerate(["=HWMAX-2", "=HWMAX-1", "=HWMAX", "=FILL"]):
    R = 6 + i
    ic.cell(R, 1, hw); ic.cell(R, 2, "=A%d/1.5" % R)
    ic.cell(R, 3, "=SQRT(G*B%d^3)" % R); ic.cell(R, 4, "=CW*A%d^1.5" % R)
    ic.cell(R, 5, "=Q/C%d" % R); ic.cell(R, 6, "=C%d/B%d" % (R, R))
    for j in range(1, 7): ic.cell(R, j).fill = CALC
ic["A12"] = "B — SUBMERGED (HDS-5): HW/D = c*(Q/(A*SQRT(D)))^2 + Y - 0.5*S0"; ic["A12"].font = SUB
for j, h in enumerate(["Entrance type", "shape", "span (ft)", "rise D (ft)", "area A (sf)",
                       "x = Q/(A*SQRT(D))", "Q per barrel (cfs)", "N for Q100",
                       "opening (ft)", "valid? x>=4"], 1):
    c = ic.cell(13, j, h); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    ic.column_dimensions[get_column_letter(j)].width = 16
ic.column_dimensions["A"].width = 32
for i in range(len(ENT)):
    R = 14 + i; SR = ENT_R0 + i
    ic.cell(R, 1, "=Inputs!A%d" % SR); ic.cell(R, 2, "=Inputs!B%d" % SR)
    ic.cell(R, 3, "=Inputs!C%d" % SR); ic.cell(R, 4, "=Inputs!D%d" % SR)
    ic.cell(R, 5, '=IF(B%d="circ",PI()*D%d^2/4,C%d*D%d)' % (R, R, R, R))
    ic.cell(R, 6, "=SQRT(MAX((HWMAX/D%d-Inputs!F%d+0.5*S0)/Inputs!E%d,0))" % (R, SR, SR))
    ic.cell(R, 7, "=F%d*E%d*SQRT(D%d)" % (R, R, R))
    ic.cell(R, 8, "=IF(G%d>0,ROUNDUP(Q/G%d,0),\"\")" % (R, R))
    ic.cell(R, 9, "=IF(N(H%d)>0,H%d*C%d,\"\")" % (R, R, R))
    ic.cell(R, 10, '=IF(F%d>=4,"yes","NO - unsubmerged, use sheet A")' % R)
    for j in range(1, 11): ic.cell(R, j).fill = CALC

# ----------------------------------------------------------- outlet control --
oc = wb.create_sheet("OutletControl")
oc["A1"] = "Outlet control  —  HW = ho + H - S0*L,  H = (1+ke+29*n^2*L/R^(4/3))*V^2/(2g)"
oc["A1"].font = Font(bold=True, size=12)
oc["A2"] = "ho comes from the Tailwater sheet. N is solved so that HW does not exceed HWmax."
oc["A2"].font = NOTE
for j, h in enumerate(["Entrance type", "span (ft)", "rise D (ft)", "area A (sf)", "R (ft)",
                       "trial N", "q per barrel", "V (ft/s)", "H (ft)", "HW outlet (ft)",
                       "HWmax (ft)", "OK?"], 1):
    c = oc.cell(4, j, h); c.font = HDR; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    oc.column_dimensions[get_column_letter(j)].width = 15
oc.column_dimensions["A"].width = 32
for i in range(len(ENT)):
    R = 5 + i; SR = ENT_R0 + i
    oc.cell(R, 1, "=Inputs!A%d" % SR); oc.cell(R, 2, "=Inputs!C%d" % SR); oc.cell(R, 3, "=Inputs!D%d" % SR)
    oc.cell(R, 4, '=IF(Inputs!B%d="circ",PI()*C%d^2/4,B%d*C%d)' % (SR, R, R, R))
    oc.cell(R, 5, '=IF(Inputs!B%d="circ",C%d/4,D%d/(2*(B%d+C%d)))' % (SR, R, R, R, R))
    c = oc.cell(R, 6, "=InletControl!H%d" % (14 + i)); c.fill = INP
    oc.cell(R, 7, "=IF(N(F%d)>0,Q/F%d,\"\")" % (R, R))
    oc.cell(R, 8, "=IF(N(F%d)>0,G%d/D%d,\"\")" % (R, R, R))
    oc.cell(R, 9, "=IF(N(F%d)>0,(1+Inputs!G%d+29*Inputs!H%d^2*L/E%d^(4/3))*H%d^2/(2*G),\"\")" % (R, SR, SR, R, R))
    oc.cell(R, 10, "=IF(N(F%d)>0,HO+I%d-S0*L,\"\")" % (R, R))
    oc.cell(R, 11, "=HWMAX")
    oc.cell(R, 12, '=IF(N(F%d)=0,"",IF(J%d<=K%d,"OK","EXCEEDS - increase N"))' % (R, R, R))
    for j in range(1, 13): oc.cell(R, j).fill = CALC
    oc.cell(R, 6).fill = INP

# ------------------------------------------------------------------ summary --
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "C10 / NM15-141 — Navajo Route 9 crossing: can a culvert pass Q100?"
sm["A1"].font = Font(bold=True, size=14)
sm.column_dimensions["A"].width = 56; sm.column_dimensions["B"].width = 16
sm.column_dimensions["C"].width = 10; sm.column_dimensions["D"].width = 58
items = [
    ("Design discharge Q100", "=Q", "cfs", "Inputs"),
    ("Available fill (road - inlet invert)", "=FILL", "ft", "the binding geometric constraint"),
    ("Allowable headwater HWmax", "=HWMAX", "ft", "fill less the freeboard allowance"),
    ("", None, None, None),
    ("TAILWATER", None, None, None),
    ("Critical depth at the outlet section", "=Tailwater!B%d" % (b + 2), "ft", "solved off the surveyed section"),
    ("Critical WSE", "=Tailwater!B%d" % (b + 3), "ft", ""),
    ("Tailwater above the INLET invert", "=TW_IN", "ft", "compare with HWmax immediately above"),
    ("Is the outlet drowned above HWmax?", '=IF(TW_IN>HWMAX,"YES - no opening size works","no")', "", "critical depth is a LOWER bound on TW"),
    ("", None, None, None),
    ("INLET CONTROL ONLY (best case, ignores tailwater)", None, None, None),
    ("Unit discharge at HWmax", "=InletControl!C8", "cfs/ft", "critical flow at the entrance"),
    ("Opening width required", "=InletControl!E8", "ft", "this is a FLOOR, not a design"),
    ("Inlet velocity at that opening", "=InletControl!F8", "ft/s", "scouring above ~10 ft/s in an unlined arroyo"),
    ("", None, None, None),
    ("EXISTING STRUCTURE", None, None, None),
]
r = 3
for lbl, f, unit, note in items:
    if f is None and lbl:
        sm.cell(r, 1, lbl).font = SUB; r += 1; continue
    if not lbl: r += 1; continue
    sm.cell(r, 1, lbl); c = sm.cell(r, 2, f); c.fill = CALC; c.border = BOX
    sm.cell(r, 3, unit); sm.cell(r, 4, note).font = NOTE
    r += 1
for i in range(EX_R1 - EX_R0 + 1):
    SR = EX_R0 + i
    sm.cell(r, 1, "=Inputs!A%d" % SR)
    sm.cell(r, 2, "=Inputs!B%d*(PI()*(Inputs!C%d/12)^2/4)*SQRT(MAX((HWMAX/(Inputs!C%d/12)-Inputs!F%d+0.5*S0)/Inputs!E%d,0))*SQRT(Inputs!C%d/12)"
            % (SR, SR, SR, ENT_R0, ENT_R0, SR))
    sm.cell(r, 3, "cfs"); sm.cell(r, 4, "=TEXT(B%d/Q,\"0.0%%\")&\" of Q100, inlet control at HWmax\"" % r).font = NOTE
    sm.cell(r, 2).fill = CALC; sm.cell(r, 2).border = BOX
    r += 1
r += 1
sm.cell(r, 1, "CONCLUSION").font = SUB; r += 1
sm.cell(r, 1, '=IF(TW_IN>HWMAX,"Outlet drowned above the allowable headwater: no culvert size passes Q100. A bridge or long multi-cell box array is required.","Head is available; see InletControl / OutletControl for the required opening.")')
sm.cell(r, 1).alignment = Alignment(wrap_text=True); sm.cell(r, 1).font = Font(bold=True)
sm.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=4)

wb.save(OUT); print("wrote", OUT)
